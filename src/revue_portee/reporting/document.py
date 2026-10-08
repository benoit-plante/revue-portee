"""A small document model shared by the Markdown and DOCX renderers.

Reports are built once as a sequence of blocks (pure functions, easy to test), then
rendered to each format. A heading may open a protocol section (``section``): the
blocks that follow, up to the next section heading, belong to it. Placeholder
paragraphs mark what the team still has to write.
"""

import io
import zipfile
from collections.abc import Iterator, Sequence
from datetime import UTC, datetime
from typing import Literal

import docx
from docx.document import Document as DocxDocument
from pydantic import AwareDatetime, BaseModel, ConfigDict, Field

__all__ = [
    "Block",
    "BulletList",
    "Code",
    "Document",
    "Heading",
    "Paragraph",
    "Table",
    "render_docx",
    "render_markdown",
    "section_blocks",
]


class Heading(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    kind: Literal["heading"] = "heading"
    level: int = Field(ge=1, le=4)
    text: str
    section: str | None = None


class Paragraph(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    kind: Literal["paragraph"] = "paragraph"
    text: str
    placeholder: bool = False


class BulletList(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    kind: Literal["bullets"] = "bullets"
    items: tuple[str, ...]


class Table(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    kind: Literal["table"] = "table"
    header: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


class Code(BaseModel):
    """Text shown as is, in a fixed-width font (search queries)."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    kind: Literal["code"] = "code"
    text: str


type Block = Heading | Paragraph | BulletList | Table | Code


class Document(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    title: str
    language: str
    generated_at: AwareDatetime
    blocks: tuple[Block, ...]


def section_blocks(blocks: Sequence[Block]) -> Iterator[tuple[str, list[Block]]]:
    """``(section, blocks)`` for every section heading, in document order.

    A section ends at the next section heading, or at a heading without section of the
    same or a higher level (e.g. "Methods" after a level-3 section): the blocks that
    follow such a heading belong to no section.
    """
    current: str | None = None
    level = 0
    content: list[Block] = []
    for block in blocks:
        starts = isinstance(block, Heading) and block.section is not None
        ends = isinstance(block, Heading) and block.level <= level
        if current is not None and (starts or ends):
            yield current, content
            current, level, content = None, 0, []
        if isinstance(block, Heading) and block.section is not None:
            current, level = block.section, block.level
        elif current is not None:
            content.append(block)
    if current is not None:
        yield current, content


def _cell(text: str) -> str:
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def render_markdown(document: Document) -> str:
    lines: list[str] = []
    for block in document.blocks:
        if isinstance(block, Heading):
            lines += ["#" * block.level + " " + block.text, ""]
        elif isinstance(block, Paragraph):
            lines += [f"_{block.text}_" if block.placeholder else block.text, ""]
        elif isinstance(block, BulletList):
            lines += [f"- {item}" for item in block.items] + [""]
        elif isinstance(block, Code):
            fence = "````" if "```" in block.text else "```"
            lines += [fence, block.text, fence, ""]
        else:
            lines.append("| " + " | ".join(_cell(h) for h in block.header) + " |")
            lines.append("|" + "---|" * len(block.header))
            lines += ["| " + " | ".join(_cell(c) for c in row) + " |" for row in block.rows]
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _docx(document: Document) -> DocxDocument:
    result = docx.Document()
    properties = result.core_properties
    properties.title = document.title
    properties.language = document.language
    properties.author = "revue-portee"
    properties.created = _naive_utc(document.generated_at)
    properties.modified = _naive_utc(document.generated_at)
    for block in document.blocks:
        if isinstance(block, Heading):
            result.add_heading(block.text, level=0 if block.level == 1 else block.level - 1)
        elif isinstance(block, Paragraph):
            paragraph = result.add_paragraph()
            run = paragraph.add_run(block.text)
            run.italic = block.placeholder
        elif isinstance(block, BulletList):
            for item in block.items:
                result.add_paragraph(item, style="List Bullet")
        elif isinstance(block, Code):
            run = result.add_paragraph().add_run(block.text)
            run.font.name = "Courier New"
        else:
            table = result.add_table(rows=1, cols=len(block.header))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, block.header, strict=True):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in block.rows:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = text
    return result


def _naive_utc(moment: datetime) -> datetime:
    # python-docx writes core property dates as UTC without time zone.
    return moment.astimezone(UTC).replace(tzinfo=None)


def _fixed_dates(content: bytes, moment: datetime) -> bytes:
    """The same package with every entry dated ``moment``: python-docx dates the entries
    of the ZIP with the current time, so two renderings of a document would differ."""
    stamp = moment.astimezone(UTC).timetuple()[:6]
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(content)) as source,
        zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for entry in source.infolist():
            fixed = zipfile.ZipInfo(entry.filename, date_time=stamp)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(fixed, source.read(entry.filename))
    return output.getvalue()


def render_docx(document: Document) -> bytes:
    """The document as DOCX; the same document always gives the same bytes."""
    buffer = io.BytesIO()
    _docx(document).save(buffer)
    return _fixed_dates(buffer.getvalue(), document.generated_at)
