"""Generated protocol: Peters et al. (2022) elements, OSF correspondence, both languages
(EF-CAD-06, EF-CAD-07, ENF-LAN-04)."""

import io
from datetime import UTC, date, datetime

import docx
import pytest

from revue_portee.domain.changes import ChangeType, CriterionChange
from revue_portee.domain.criteria import (
    CriteriaVersion,
    Criterion,
    CriterionKind,
    PccElement,
    VersionStatus,
)
from revue_portee.domain.framing import Framing
from revue_portee.domain.project import Project, Reviewer, ReviewerKind
from revue_portee.domain.protocol import (
    FREE_TEXT_SECTIONS,
    ProtocolRegistration,
    ProtocolSection,
    ProtocolText,
)
from revue_portee.reporting.document import (
    Block,
    BulletList,
    Document,
    Heading,
    Paragraph,
    Table,
    render_docx,
    render_markdown,
    section_blocks,
)
from revue_portee.reporting.protocol import (
    Deviation,
    GridDeviation,
    ProtocolData,
    build_protocol,
    checklist_status,
)
from revue_portee.resources import default_ai_settings, osf_form, peters_checklist

NOW = datetime(2026, 10, 7, 12, 0, tzinfo=UTC)
PROJECT = Project(id="p", title="Soutien à la parentalité", language="fr", created_at=NOW)
TEAM = (Reviewer(id="h", kind=ReviewerKind.HUMAN, display_name="Benoit Plante"),)
CRITERIA = CriteriaVersion(
    id="v2",
    number=2,
    parent_id="v1",
    status=VersionStatus.ACTIVE,
    created_at=NOW,
    activated_at=NOW,
    author_id="h",
    rationale="Adolescents",
    criteria=(
        Criterion(
            code="P1",
            pcc_element=PccElement.POPULATION,
            kind=CriterionKind.INCLUSION,
            text="Parents d'enfants de 0 à 17 ans",
            examples=("Pères", "Mères"),
        ),
        Criterion(
            code="C1",
            pcc_element=PccElement.CONCEPT,
            kind=CriterionKind.INCLUSION,
            text="Intervention | soutien",
            guidance="Programme structuré",
        ),
        Criterion(
            code="CTX1",
            pcc_element=PccElement.CONTEXT,
            kind=CriterionKind.INCLUSION,
            text="Services communautaires",
        ),
        Criterion(
            code="X1",
            pcc_element=PccElement.OTHER,
            kind=CriterionKind.EXCLUSION,
            text="Éditoriaux",
        ),
    ),
)
FULL_TEXT = ProtocolText(sections={s: f"Texte {s.value}." for s in FREE_TEXT_SECTIONS})


def data(**changes: object) -> ProtocolData:
    values: dict[str, object] = {
        "project": PROJECT,
        "reviewers": TEAM,
        "framing": None,
        "criteria": None,
        "text": ProtocolText(),
        "ai": default_ai_settings(),
        "registration": None,
        "tool_version": "0.0.0-test",
        "generated_at": NOW,
    }
    return ProtocolData.model_validate(values | changes)


FULL = data(
    framing=Framing(
        question="Quelles interventions ?",
        population="Parents",
        concept="Soutien",
        context="Communautaire",
        secondary_questions=("Quels résultats ?",),
    ),
    criteria=CRITERIA,
    text=FULL_TEXT,
)


def build(value: ProtocolData, language: str = "fr") -> Document:
    return build_protocol(value, language=language, checklist=peters_checklist(), osf=osf_form())


@pytest.mark.parametrize("language", ["fr", "en"])
def test_every_peters_element_has_its_section(language: str) -> None:
    for value in (data(), FULL):
        statuses = checklist_status(build(value, language).blocks, peters_checklist())
        assert len(statuses) == len(peters_checklist().items) == 23
        assert all(s.present for s in statuses)


def test_a_complete_project_fills_every_element() -> None:
    statuses = checklist_status(build(FULL).blocks, peters_checklist())
    assert [s.item_id for s in statuses if not s.filled] == []


def test_an_empty_project_shows_what_is_missing() -> None:
    statuses = {s.item_id: s for s in checklist_status(build(data()).blocks, peters_checklist())}
    # Generated from the project and the AI configuration alone:
    filled = {k for k, s in statuses.items() if s.filled}
    assert filled == {"PETERS-01", "PETERS-11", "PETERS-14", "PETERS-22", "PETERS-23"}
    text = render_markdown(build(data()))
    assert "_À compléter : question principale (page « Cadrage »)_" in text


def test_french_and_english_versions() -> None:
    french = render_markdown(build(FULL, "fr"))
    english = render_markdown(build(FULL, "en"))
    assert french.startswith("# Soutien à la parentalité : protocole de revue de portée\n")
    assert english.startswith("# Soutien à la parentalité: a scoping review protocol\n")
    assert "## Méthode" in french
    assert "## Methods" in english
    # The team's own text is reproduced as written in both languages.
    assert "Texte background." in french
    assert "Texte background." in english
    assert "- P1 (inclusion) : Parents d'enfants de 0 à 17 ans" in french
    assert "- X1 (exclusion): Éditoriaux" in english


def test_ai_section_comes_from_the_project_configuration() -> None:
    settings = default_ai_settings()
    text = render_markdown(build(data()))
    model = settings.enabled_task("suggest_pcc").model
    assert (
        f"| Suggestions pour le cadrage (reformulations, éléments PCC) | anthropic — {model} |"
        in text
    )
    screener = settings.enabled_task("screen_reference").model
    assert f"| Second réviseur au tri des titres et résumés | anthropic — {screener} |" in text
    assert f"échantillon aléatoire de {settings.supervision.pilot_sample_size} références" in text
    target = str(settings.supervision.target_sensitivity)
    assert f"au moins {target.replace('.', ',')}." in text
    assert f"at least {target}." in render_markdown(build(data(), "en"))
    assert "n'exclut jamais une référence à elle seule" in text


def test_registration_and_deviations() -> None:
    registration = ProtocolRegistration(
        id="r",
        doi="10.17605/OSF.IO/ABCDE",
        registered_on=date(2026, 10, 1),
        criteria_version_id="v1",
        created_at=NOW,
        reviewer_id="h",
    )
    change = CriterionChange(
        id="c",
        from_version_id="v1",
        to_version_id="v2",
        code="P1",
        change_type=ChangeType.BROADENING,
        proposed_by=ReviewerKind.AI,
        confirmed_by="h",
        created_at=NOW,
    )
    deviation = Deviation(number=2, activated_at=NOW, rationale="Adolescents", changes=(change,))
    text = render_markdown(build(data(registration=registration, deviations=(deviation,))))
    assert "Enregistrement : DOI 10.17605/OSF.IO/ABCDE, enregistré le 2026-10-01." in text
    assert "Le protocole est enregistré sur OSF : DOI 10.17605/OSF.IO/ABCDE." in text
    assert "Version 2 des critères (2026-10-07) : Adolescents — P1 (élargissement)" in text
    grid = GridDeviation(
        number=2, activated_at=NOW, rationale="Taille définie", added=("D4",),
        modified=("D2",), removed=("D3",),
    )  # fmt: skip
    with_grid = render_markdown(build(data(registration=registration, grid_deviations=(grid,))))
    assert (
        "Version 2 de la grille d'extraction (2026-10-07) : Taille définie — D4 (ajout), "
        "D2 (modification), D3 (retrait)"
    ) in with_grid
    unregistered = render_markdown(build(data()))
    assert "pas encore enregistré" in unregistered


def test_osf_correspondence_covers_the_65_items() -> None:
    document = build(FULL, "en")
    tables = [b for b in document.blocks if isinstance(b, Table) and b.header[0] == "OSF item"]
    (table,) = tables
    assert len(table.rows) == 65
    rows = {row[0]: row for row in table.rows}
    assert rows["GSRRF-11"] == ("GSRRF-11", "Primary research question(s)", "Review question")
    assert rows["GSRRF-24"][2] == ("Participants; Concept; Context; Types of sources of evidence")
    assert rows["GSRRF-8"][2] == "to be filled in directly in the OSF form, or not applicable"


def test_peters_appendix_reports_the_status_of_each_element() -> None:
    document = build(data(), "en")
    (table,) = [b for b in document.blocks if isinstance(b, Table) and b.header[0] == "Element"]
    rows = {row[0]: row for row in table.rows}
    assert rows["PETERS-01"][3] == "written"
    assert rows["PETERS-03"] == (
        "PETERS-03",
        "Introduction - rationale and what is already known",
        "Background and rationale",
        "to be completed",
    )
    assert rows["PETERS-18"][3] == "to be completed if applicable"
    text = render_markdown(document)
    assert "to be checked against the checklist of the article" in text


def test_criteria_appendix_and_markdown_escaping() -> None:
    text = render_markdown(build(FULL))
    assert (
        "| C1 | Concept | inclusion | Intervention \\| soutien | Programme structuré |  |  |"
        in text
    )
    assert (
        "| P1 | Population | inclusion | Parents d'enfants de 0 à 17 ans |  | Pères; Mères |  |"
        in text
    )


def test_markdown_rendering_of_each_block() -> None:
    document = Document(
        title="t",
        language="fr",
        generated_at=NOW,
        blocks=(
            Heading(level=2, text="Titre"),
            Paragraph(text="Texte"),
            Paragraph(text="Manque", placeholder=True),
            BulletList(items=("a", "b")),
            Table(header=("A", "B"), rows=(("1", "x\ny"),)),
        ),
    )
    assert render_markdown(document) == (
        "## Titre\n\nTexte\n\n_Manque_\n\n- a\n- b\n\n| A | B |\n|---|---|\n| 1 | x<br>y |\n"
    )


def test_docx_rendering() -> None:
    document = build(FULL, "fr")
    parsed = docx.Document(io.BytesIO(render_docx(document)))
    assert parsed.core_properties.title == document.title
    assert parsed.core_properties.language == "fr"
    assert parsed.core_properties.created == NOW
    headings = [
        p.text
        for p in parsed.paragraphs
        if p.style is not None and str(p.style.name).startswith(("Heading", "Title"))
    ]
    expected = [b.text for b in document.blocks if isinstance(b, Heading)]
    assert headings == expected
    tables = [b for b in document.blocks if isinstance(b, Table)]
    assert len(parsed.tables) == len(tables)
    first = parsed.tables[0]
    assert [c.text for c in first.rows[0].cells] == list(tables[0].header)
    assert len(first.rows) == len(tables[0].rows) + 1


def test_section_blocks() -> None:
    blocks: list[Block] = [
        Paragraph(text="before"),
        Heading(level=2, text="A", section="a"),
        Paragraph(text="1"),
        Heading(level=3, text="sub"),
        Heading(level=2, text="B", section="b"),
    ]
    assert list(section_blocks(blocks)) == [("a", [blocks[2], blocks[3]]), ("b", [])]
    assert list(section_blocks([])) == []


def test_sections_follow_the_document_order() -> None:
    document = build(FULL)
    order = [s for s, _ in section_blocks(list(document.blocks))]
    assert order == [s.value for s in ProtocolSection]


def test_text_under_a_plain_heading_belongs_to_no_section() -> None:
    # Criteria in force but no framing: the "Criteria version…" paragraph under the
    # "Eligibility criteria" heading must not fill the review question.
    statuses = {
        s.item_id: s
        for s in checklist_status(build(data(criteria=CRITERIA)).blocks, peters_checklist())
    }
    assert not statuses["PETERS-06"].filled
    assert statuses["PETERS-07"].filled  # P1 is listed under Participants
